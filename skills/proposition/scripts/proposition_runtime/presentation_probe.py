"""Optional local-tool fixture, never a client deck or an E2E quality verdict."""
import importlib.metadata
from pathlib import Path
import shutil
import subprocess
import tempfile


def inventory():
    try:
        version = importlib.metadata.version('python-pptx')
    except importlib.metadata.PackageNotFoundError:
        version = None
    return {'python_pptx': version, 'renderer': shutil.which('libreoffice') or shutil.which('soffice')}


def probe():
    tools = inventory()
    result = {'status': 'BLOCKED', 'scope': 'local_editable_fixture_only', 'tools': tools,
              'authoring': 'NOT_RUN', 'editing': 'NOT_RUN', 'rendering': 'NOT_RUN',
              'rasterization': 'NOT_RUN', 'client_fidelity': 'NOT_RUN'}
    if not tools['python_pptx']:
        result['authoring'] = 'UNAVAILABLE'
    if not tools['renderer']:
        result['rendering'] = 'UNAVAILABLE'
    if not tools['python_pptx'] or not tools['renderer']:
        return result
    try:
        from PIL import Image
        from pptx import Presentation
        from pptx.util import Inches
        from pypdf import PdfReader
        import pypdfium2
        with tempfile.TemporaryDirectory(prefix='proposition-pptx-probe-') as folder:
            root = Path(folder)
            image = root / 'fixture.png'
            Image.new('RGB', (160, 90), '#305070').save(image)
            document = Presentation()
            slide = document.slides.add_slide(document.slide_layouts[6])
            slide.shapes.add_textbox(Inches(1), Inches(1), Inches(7), Inches(1)).text = 'Original fixture'
            slide.shapes.add_picture(str(image), Inches(1), Inches(3), width=Inches(3))
            initial = root / 'initial.pptx'
            document.save(initial)
            result['authoring'] = 'PASS'
            edited = Presentation(initial)
            edited.slides[0].shapes[0].text = 'Edited fixture 42'
            final = root / 'edited.pptx'
            edited.save(final)
            reopened = Presentation(final)
            if reopened.slides[0].shapes[0].text != 'Edited fixture 42' or len(reopened.slides[0].shapes) != 2:
                raise ValueError('Editable text or separate image did not survive the round trip')
            result['editing'] = 'PASS'
            rendered = subprocess.run([tools['renderer'], '--headless',
                '-env:UserInstallation=' + (root / 'office-profile').as_uri(),
                '--convert-to', 'pdf', '--outdir', str(root), str(final)],
                capture_output=True,text=True,timeout=90)
            pdf = root / 'edited.pdf'
            if rendered.returncode != 0 or not pdf.is_file():
                raise ValueError('Renderer did not produce the fixture PDF')
            reader = PdfReader(pdf)
            if len(reader.pages) != 1 or 'Edited fixture 42' not in reader.pages[0].extract_text():
                raise ValueError('Rendered PDF does not contain the edited slide')
            result.update(rendering='PASS',page_count=1)
            raster = pypdfium2.PdfDocument(str(pdf))
            try:
                page = raster[0]
                try:
                    bitmap = page.render(scale=1)
                    try:
                        pixels = bitmap.to_pil()
                        pixels.save(root / 'render.png')
                        if pixels.width < 100 or pixels.height < 100:
                            raise ValueError('Rendered slide is too small to inspect')
                    finally:
                        bitmap.close()
                finally:
                    page.close()
            finally:
                raster.close()
            result.update(rasterization='PASS',status='PASS')
    except Exception as error:
        result['error'] = str(error)
    return result
